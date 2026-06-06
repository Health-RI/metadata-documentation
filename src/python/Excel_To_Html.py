from pathlib import Path
import pandas as pd
import re
import json
from openpyxl import load_workbook
from urllib.parse import urlparse, urlunparse


EXCEL_FILE_PATH = "../excel/HealthRI_v2.0.2.xlsx"
OUTPUT_PATH = Path("../property")
LINKS_FILE = Path(__file__).parent / "links.json"


# ========================
# URL HELPERS
# ========================

def extract_urls(text):
    if not isinstance(text, str):
        return []
    return re.findall(r'https?://[^\s\]\),]+', text)


def strip_urls(text):
    if not isinstance(text, str):
        return ""
    return re.sub(r'https?://[^\s\]\),]+', '', text)


def make_clickable(url):
    return f'<a href="{url}">{url}</a>'


def normalize_url(url):
    parsed = urlparse(url)
    return urlunparse((
        "",
        parsed.netloc,
        parsed.path.rstrip("/"),
        "", "", ""
    ))


def select_canonical_urls(urls):
    groups = {}

    for url in urls:
        if not url:
            continue

        base = normalize_url(url)
        groups.setdefault(base, []).append(url)

    final = []

    for variants in groups.values():
        fragments = [u for u in variants if "#" in u]
        if fragments:
            final.append(fragments[0])
        else:
            https = [u for u in variants if u.startswith("https")]
            final.append(https[0] if https else variants[0])

    return final


# ========================
# LABEL → URL MAPPING
# ========================

def resolve_vocab_labels(text, mapping):
    if not isinstance(text, str):
        return []

    return [
        url for label, url in mapping.items()
        if label.lower() in text.lower()
    ]


# ========================
# REQUIREMENT DETECTION (fallback only)
# ========================

def detect_requirement(text):
    if not isinstance(text, str):
        return "MAY"

    t = text.lower()

    if re.search(r"at\s+least", t):
        return "AT_LEAST_1"

    # ✅ tightening this prevents false MUST
    if "must " in t:
        return "MUST"

    return "MAY"


# ========================
# EXCEL URL EXTRACTION
# ========================

def extract_urls_per_property(sheet_name):
    wb = load_workbook(EXCEL_FILE_PATH)
    ws = wb[sheet_name]

    headers = [c.value for c in ws[1]]
    vocab_idx = headers.index("Controlled vocabluary (if applicable)") + 1
    uri_idx = headers.index("Property URI") + 1

    url_map = {}

    for row_idx in range(2, ws.max_row + 1):
        uri = ws.cell(row=row_idx, column=uri_idx).value
        cell = ws.cell(row=row_idx, column=vocab_idx)

        if not uri:
            continue

        urls = []

        if cell.value:
            urls += extract_urls(str(cell.value))

        if cell.hyperlink:
            urls.append(cell.hyperlink.target)

        url_map[uri] = urls

    return url_map


# ========================
# MERGE MULTI-ROW VOCAB
# ========================

def merge_vocab_rows(df):
    merged = []
    current = None

    for _, row in df.iterrows():
        vocab = row["Controlled vocabluary (if applicable)"]

        is_cont = (
            pd.notna(vocab)
            and all(
                pd.isna(v) or str(v).strip() == ""
                for col, v in row.items()
                if col != "Controlled vocabluary (if applicable)"
            )
        )

        if is_cont and current is not None:
            existing = current["Controlled vocabluary (if applicable)"]
            current["Controlled vocabluary (if applicable)"] = (
                str(existing) + "\n" + str(vocab) if pd.notna(existing) else str(vocab)
            )
        else:
            if current is not None:
                merged.append(current)
            current = row.copy()

    if current is not None:
        merged.append(current)

    return pd.DataFrame(merged).reindex(columns=df.columns)


# ========================
# ✅ FINAL USAGE NOTE (CORRECT + SHEET-PROOF)
# ========================

def build_usage_note(row, sheet_name, url_map, links):

    base = row["Usage note"]
    if not isinstance(base, str) or not base.strip():
        base = "N.A."

    vocab = row["Controlled vocabluary (if applicable)"]
    prop = row["Property URI"]

    if not isinstance(vocab, str) or not vocab.strip():
        return base

    key_exact = f"{prop}|{sheet_name}"
    key_lower = f"{prop}|{sheet_name.lower()}"

    prop_key_norm = prop.strip().lower()

    # ✅ ✅ FINAL FIXED LOOKUP
    vocab_type = (
        links["vocabRequirements"].get(key_exact)
        or links["vocabRequirements"].get(key_lower)
        or next(
            (
                v
                for k, v in links["vocabRequirements"].items()
                if k.split("|")[0].strip().lower() == prop_key_norm
            ),
            None
        )
        or detect_requirement(vocab)
    )

    vocab_text = links["vocabTexts"].get(vocab_type, "")

    # URLs
    text_urls = extract_urls(vocab)
    excel_urls = url_map.get(prop, [])
    mapped_urls = resolve_vocab_labels(vocab, links["vocabLabelMapping"])

    urls = select_canonical_urls(text_urls + excel_urls + mapped_urls)

    # clean vocab text
    mapping_labels = [k.lower() for k in links["vocabLabelMapping"].keys()]
    clean_lines = []

    for line in strip_urls(vocab).split("\n"):
        l = line.strip()
        if not l:
            continue

        lower = l.lower()

        if any(label in lower for label in mapping_labels):
            continue

        if len(l.split()) <= 6:
            continue

        if not re.search(r"\b(is|are|must|should|may|use|used|provides|represents)\b", lower):
            continue

        clean_lines.append(l)

    clean_text = "<br>".join(clean_lines)

    parts = []

    if vocab_text:
        parts.append(vocab_text)

    if clean_text:
        parts.append(clean_text)

    if urls:
        parts.append("")
        parts.extend([make_clickable(u) for u in urls])
        parts.append("")

    vocab_block = "<br>".join(parts)

    return base + "<br><br><strong>Controlled vocabulary</strong><br>" + vocab_block


# ========================
# MAIN
# ========================

def main():
    OUTPUT_PATH.mkdir(parents=True, exist_ok=True)

    classes_df = pd.read_excel(EXCEL_FILE_PATH, sheet_name="classes")

    with open(LINKS_FILE, "r", encoding="utf-8") as f:
        links = json.load(f)

    for _, class_row in classes_df.iterrows():
        sheet_name = class_row["sheet_name"]

        df = pd.read_excel(EXCEL_FILE_PATH, sheet_name=sheet_name)

        cols = [
            "Property label",
            "Definition",
            "Property URI",
            "Range",
            "Cardinality",
            "Usage note",
            "Controlled vocabluary (if applicable)"
        ]

        df = df[cols]
        df = merge_vocab_rows(df)

        url_map = extract_urls_per_property(sheet_name)

        df["Usage note"] = df.apply(
            lambda r: build_usage_note(r, sheet_name, url_map, links),
            axis=1
        )

        df = df.drop(columns=["Controlled vocabluary (if applicable)"])

        output_file = OUTPUT_PATH / f"properties-{sheet_name.lower()}.html"
        df.to_html(output_file, index=False, escape=False)

        print(f"Generated {output_file}")


if __name__ == "__main__":
    main()